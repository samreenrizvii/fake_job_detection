from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

def create_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "TrueJob: Fake Job Posting Detection"
    subtitle.text = "A Machine Learning approach using NLP\nPresented by: Arman Rizvi\n3rd Year B.Tech, Data Science & AI"

    # Define a helper function to add standard slides
    def add_slide(title_text, bullet_points):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        title.text = title_text
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.text = bullet_points[0]
        for point in bullet_points[1:]:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
            
    # Slide 2: Problem Statement
    add_slide("Problem Statement", [
        "Online job portals are flooded with fraudulent listings.",
        "Scams target vulnerable job seekers to extract money or personal info.",
        "Traditional spam filters fail because scams use professional, legitimate-sounding language.",
        "There is a need for a real-time, automated detection tool."
    ])

    # Slide 3: Objectives
    add_slide("Project Objectives", [
        "Preprocess and extract features from unstructured job posting text.",
        "Train and compare multiple Machine Learning algorithms.",
        "Address severe class imbalance mathematically.",
        "Develop an automated pipeline for intelligent model selection.",
        "Build an interactive web dashboard for real-time predictions."
    ])

    # Slide 4: The Dataset (EMSCAD)
    add_slide("Dataset: EMSCAD", [
        "Employment Scam Aegean Dataset (from Kaggle).",
        "Total Records: ~17,880 job postings.",
        "Features used: Title, Company Profile, Description, Requirements, Benefits.",
        "Target Variable: 'fraudulent' (0 = Real, 1 = Fake).",
        "Challenge: Severe Imbalance. Only ~5% of jobs are fraudulent."
    ])

    # Slide 5: Data Preprocessing & NLP
    add_slide("Data Preprocessing & NLP", [
        "Missing Data: Filled empty text fields instead of dropping rows.",
        "Feature Engineering: Concatenated all text features into a single 'full_text' column.",
        "NLP Cleaning: Converted to lowercase, removed punctuation and numbers.",
        "Vectorization: Used TF-IDF (Term Frequency-Inverse Document Frequency).",
        "N-grams: Used unigrams and bigrams (e.g., 'wire transfer') to capture context."
    ])

    # Slide 6: Machine Learning Algorithms
    add_slide("Algorithms Evaluated", [
        "1. Logistic Regression: Fast, interpretable baseline.",
        "2. Naive Bayes: Classic NLP text classifier.",
        "3. Random Forest: Ensemble method to capture non-linear patterns.",
        "4. Linear SVC: Support Vector Machine optimized for high-dimensional, sparse TF-IDF matrices.",
        "Crucial Step: Used 'class_weight=balanced' to force algorithms to penalize errors on minority class."
    ])

    # Slide 7: Evaluation Metrics
    add_slide("Why Accuracy is Misleading", [
        "Because 95% of the data is Real, a 'dumb' model predicting 'Real' every time achieves 95% Accuracy.",
        "Therefore, we prioritized different metrics:",
        "Recall: Out of all actual scams, how many did we catch?",
        "Precision: Out of all flagged jobs, how many were actually scams?",
        "F1-Score: The harmonic mean of Precision and Recall. This was our primary metric for model selection."
    ])

    # Slide 8: Training Results
    add_slide("Model Performance & Selection", [
        "Logistic Regression: Accuracy 96.7% | F1-Score 72.8%",
        "Naive Bayes: Accuracy 97.0% | F1-Score 58.4%",
        "Random Forest: Accuracy 98.1% | F1-Score 77.8%",
        "Linear SVC: Accuracy 98.1% | F1-Score 81.5% | Recall 82.6%",
        "Winner: Linear SVC. Automatically selected by the pipeline for achieving the highest F1-Score."
    ])

    # Slide 9: Dash Web Application
    add_slide("Interactive Dashboard", [
        "Built using Python Dash and Dash Bootstrap Components.",
        "Features a 'Cyborg' dark-mode premium UI.",
        "Real-time Inference: Users paste a job description, and the backend vectorizes and predicts instantly.",
        "Model Insights: Visualizes exactly which words the model flagged as 'Fraudulent' based on SVC coefficients."
    ])

    # Slide 10: Conclusion & Future Scope
    add_slide("Conclusion & Future Scope", [
        "Conclusion:",
        "- Successfully built an end-to-end ML pipeline.",
        "- Proved that Linear SVC outperforms others on sparse text data.",
        "- Demonstrated the danger of relying purely on Accuracy.",
        "Future Scope:",
        "- Upgrade from TF-IDF to Transformer models (BERT).",
        "- Deploy as a Google Chrome browser extension for live web scanning."
    ])

    # Save the presentation
    save_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'docs', 'TrueJob_Presentation.pptx')
    prs.save(save_path)
    print(f"Presentation generated successfully at: {save_path}")

if __name__ == "__main__":
    create_presentation()
